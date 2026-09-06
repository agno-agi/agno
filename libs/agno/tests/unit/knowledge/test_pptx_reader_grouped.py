"""
Regression test for agno#9999:
PPTXReader must extract text from shapes inside GroupShapes (including nested groups).
"""
import pytest
from io import BytesIO

from pptx import Presentation
from pptx.util import Inches


def _make_pptx(grouped: bool) -> BytesIO:
    """Create a minimal PPTX with one text box, optionally inside a group."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text = "Launch date: 2026-10-01"
    if grouped:
        slide.shapes.add_group_shape([box])
    buf = BytesIO()
    deck.save(buf)
    return buf


def _make_nested_pptx() -> BytesIO:
    """Create a PPTX with a text box nested two levels deep in groups."""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text = "Nested text"
    outer = slide.shapes.add_group_shape([box])
    slide.shapes.add_group_shape([outer])
    buf = BytesIO()
    deck.save(buf)
    return buf


def test_grouped_text_is_extracted():
    """Text inside a GroupShape must appear in the extracted content."""
    from agno.knowledge.reader.pptx_reader import PPTXReader

    reader = PPTXReader(chunk=False)
    plain_docs = reader.read(BytesIO(_make_pptx(grouped=False).getvalue()))
    grouped_docs = reader.read(BytesIO(_make_pptx(grouped=True).getvalue()))

    assert plain_docs, "plain PPTX returned no documents"
    assert grouped_docs, "grouped PPTX returned no documents"

    plain_text = plain_docs[0].content
    grouped_text = grouped_docs[0].content

    assert "Launch date: 2026-10-01" in plain_text, "plain text not found"
    assert "Launch date: 2026-10-01" in grouped_text, (
        f"Grouped text was dropped. Got: {repr(grouped_text)}"
    )
    assert "(No text content)" not in grouped_text, (
        "Reader reported no text content for a slide that has grouped text"
    )


def test_nested_group_text_is_extracted():
    """Text nested two levels deep in GroupShapes must also be extracted."""
    from agno.knowledge.reader.pptx_reader import PPTXReader

    reader = PPTXReader(chunk=False)
    docs = reader.read(BytesIO(_make_nested_pptx().getvalue()))

    assert docs, "nested PPTX returned no documents"
    assert "Nested text" in docs[0].content, (
        f"Doubly-nested text was dropped. Got: {repr(docs[0].content)}"
    )


def test_ungrouped_shapes_still_work():
    """Existing behaviour for plain (non-grouped) shapes must be preserved."""
    from agno.knowledge.reader.pptx_reader import PPTXReader

    reader = PPTXReader(chunk=False)
    docs = reader.read(BytesIO(_make_pptx(grouped=False).getvalue()))
    assert docs and "Launch date: 2026-10-01" in docs[0].content
