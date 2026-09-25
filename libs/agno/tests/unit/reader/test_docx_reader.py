import asyncio
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

import pytest
from docx import Document as RealDocument

from agno.knowledge.document.base import Document
from agno.knowledge.reader.docx_reader import DocxReader


@pytest.fixture
def mock_docx():
    """Provide a DOCX document with real paragraph and XML interfaces."""
    mock_doc = RealDocument()
    mock_doc.add_paragraph("First paragraph")
    mock_doc.add_paragraph("Second paragraph")
    return mock_doc


def test_docx_reader_read_file(mock_docx):
    """Test reading a DOCX file"""
    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_docx),
    ):
        reader = DocxReader()
        documents = reader.read(Path("test.docx"))

        assert len(documents) == 1
        assert documents[0].name == "test"
        assert documents[0].content == "First paragraph\n\nSecond paragraph"


@pytest.mark.asyncio
async def test_docx_reader_async_read_file(mock_docx):
    """Test reading a DOCX file asynchronously"""
    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_docx),
    ):
        reader = DocxReader()
        documents = await reader.async_read(Path("test.docx"))

        assert len(documents) == 1
        assert documents[0].name == "test"
        assert documents[0].content == "First paragraph\n\nSecond paragraph"


def test_docx_reader_with_chunking():
    """Test reading a DOCX file with chunking enabled"""
    mock_doc = RealDocument()
    mock_doc.add_paragraph("Test content")

    chunked_docs = [
        Document(name="test", id="test_1", content="Chunk 1"),
        Document(name="test", id="test_2", content="Chunk 2"),
    ]

    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_doc),
    ):
        reader = DocxReader()
        reader.chunk = True
        reader.chunk_document = Mock(return_value=chunked_docs)

        documents = reader.read(Path("test.docx"))

        reader.chunk_document.assert_called_once()
        assert len(documents) == 2
        assert documents[0].content == "Chunk 1"
        assert documents[1].content == "Chunk 2"


def test_docx_reader_bytesio(mock_docx):
    """Test reading a DOCX from BytesIO"""
    file_obj = BytesIO(b"dummy content")
    file_obj.name = "test.docx"

    with patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_docx):
        reader = DocxReader()
        documents = reader.read(file_obj)

        assert len(documents) == 1
        assert documents[0].name == "test"
        assert documents[0].content == "First paragraph\n\nSecond paragraph"


def test_docx_reader_invalid_file():
    """Test reading an invalid file"""
    with patch("pathlib.Path.exists", return_value=False):
        reader = DocxReader()
        documents = reader.read(Path("nonexistent.docx"))
        assert len(documents) == 0


def test_docx_reader_file_error():
    """Test handling of file reading errors"""
    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", side_effect=Exception("File error")),
    ):
        reader = DocxReader()
        documents = reader.read(Path("test.docx"))
        assert len(documents) == 0


@pytest.mark.asyncio
async def test_async_docx_processing(mock_docx):
    """Test concurrent async processing"""
    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_docx),
    ):
        reader = DocxReader()
        tasks = [reader.async_read(Path("test.docx")) for _ in range(3)]
        results = await asyncio.gather(*tasks)

        assert len(results) == 3
        assert all(len(docs) == 1 for docs in results)
        assert all(docs[0].name == "test" for docs in results)
        assert all(docs[0].content == "First paragraph\n\nSecond paragraph" for docs in results)


@pytest.mark.asyncio
async def test_docx_reader_async_with_chunking():
    """Test async reading with chunking enabled"""
    mock_doc = RealDocument()
    mock_doc.add_paragraph("Test content")

    # Create a chunked document
    chunked_docs = [
        Document(name="test", id="test_1", content="Chunk 1"),
        Document(name="test", id="test_2", content="Chunk 2"),
    ]

    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_doc),
    ):
        reader = DocxReader()
        reader.chunk = True
        # Mock the chunk_document method to return our predefined chunks
        reader.chunk_document = Mock(return_value=chunked_docs)

        documents = await reader.async_read(Path("test.docx"))

        reader.chunk_document.assert_called_once()
        assert len(documents) == 2
        assert documents[0].content == "Chunk 1"
        assert documents[1].content == "Chunk 2"


def test_docx_reader_metadata(mock_docx):
    """Test document metadata"""
    with (
        patch("pathlib.Path.exists", return_value=True),
        patch("agno.knowledge.reader.docx_reader.DocxDocument", return_value=mock_docx),
    ):
        reader = DocxReader()
        documents = reader.read(Path("test_doc.docx"))

        assert len(documents) == 1
        assert documents[0].name == "test_doc"
        assert documents[0].content == "First paragraph\n\nSecond paragraph"


def test_docx_reader_chunk_size_propagation():
    """Test that chunk_size is propagated to default chunking strategy"""
    from agno.knowledge.chunking.document import DocumentChunking

    reader = DocxReader(chunk_size=350)
    assert reader.chunk_size == 350
    assert reader.chunking_strategy.chunk_size == 350
    assert isinstance(reader.chunking_strategy, DocumentChunking)


def test_docx_reader_default_chunk_size():
    """Test default chunk_size is 5000"""
    from agno.knowledge.chunking.document import DocumentChunking

    reader = DocxReader()
    assert reader.chunk_size == 5000
    assert reader.chunking_strategy.chunk_size == 5000
    assert isinstance(reader.chunking_strategy, DocumentChunking)


def test_docx_reader_reads_tables():
    """Test reading a real DOCX file with tables and paragraphs"""
    doc = RealDocument()
    doc.add_paragraph("Before table")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "Agno"
    table.cell(1, 1).text = "100"
    doc.add_paragraph("After table")

    source = BytesIO()
    doc.save(source)
    source.seek(0)

    reader = DocxReader(chunk=False)
    documents = reader.read(source)
    assert len(documents) == 1
    content = documents[0].content
    assert "Before table" in content
    assert "Product | Revenue" in content
    assert "Agno | 100" in content
    assert "After table" in content
