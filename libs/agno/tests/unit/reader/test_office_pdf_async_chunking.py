"""Exercise awaited chunking through real Office/PDF parsing."""

import asyncio

import pytest

pytest.importorskip("docx")
pytest.importorskip("pptx")
pytest.importorskip("pypdf")
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from agno.knowledge.chunking.strategy import ChunkingStrategy
from agno.knowledge.document.base import Document
from agno.knowledge.reader.docx_reader import DocxReader
from agno.knowledge.reader.pdf_reader import PDFReader
from agno.knowledge.reader.pptx_reader import PPTXReader


class AwaitedChunking(ChunkingStrategy):
    """A supported async extension that must execute on the caller's loop."""

    def __init__(self):
        self.loop = asyncio.get_running_loop()
        self.inputs = []

    def chunk(self, document):
        raise RuntimeError("Use awaited chunking")

    async def achunk(self, document):
        assert asyncio.get_running_loop() is self.loop
        await asyncio.sleep(0)
        self.inputs.append(document)
        return [Document(name=document.name, id=document.id, meta_data=document.meta_data, content="Awaited chunk")]


@pytest.fixture(params=["docx", "pptx", "pdf"])
def parsed_file(request, tmp_path):
    path = tmp_path / f"content.{request.param}"
    if request.param == "docx":
        package = WordDocument()
        package.add_paragraph("Native parsed content")
        package.save(path)
        reader_type = DocxReader
    elif request.param == "pptx":
        package = Presentation()
        slide = package.slides.add_slide(package.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Native parsed content"
        package.save(path)
        reader_type = PPTXReader
    else:
        package = PdfWriter()
        page = package.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): package._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 200 Td (Native parsed content) Tj ET")
        page[NameObject("/Contents")] = package._add_object(content)
        package.write(path)
        reader_type = PDFReader
    return path, reader_type


async def test_async_read_awaits_configured_chunking(parsed_file):
    path, reader_type = parsed_file
    strategy = AwaitedChunking()
    reader = reader_type(chunking_strategy=strategy)
    documents = await asyncio.wait_for(reader.async_read(path), timeout=5)
    assert len(strategy.inputs) == 1
    assert "Native parsed content" in strategy.inputs[0].content
    assert [document.content for document in documents] == ["Awaited chunk"]
    assert documents[0].name == "content"
    assert reader.chunk is True


async def test_async_read_without_chunking_keeps_native_content(parsed_file):
    path, reader_type = parsed_file
    strategy = AwaitedChunking()
    reader = reader_type(chunk=False, chunking_strategy=strategy)
    documents = await asyncio.wait_for(reader.async_read(path), timeout=5)
    assert len(documents) == 1
    assert "Native parsed content" in documents[0].content
    assert strategy.inputs == []
    assert reader.chunk is False
