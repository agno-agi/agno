"""Real Office packages uploaded through native temporary file streams."""

import io
import tempfile

import pytest

pytest.importorskip("docx")
pytest.importorskip("pptx")
from docx import Document as WordDocument
from pptx import Presentation
from pptx.util import Inches

from agno.knowledge.reader.docx_reader import DocxReader
from agno.knowledge.reader.pptx_reader import PPTXReader


@pytest.fixture(params=["docx", "pptx"])
def office_package(request):
    buffer = io.BytesIO()
    if request.param == "docx":
        package = WordDocument()
        package.add_paragraph("Uploaded office content")
        reader = DocxReader(chunk=False)
        expected_name = "docx_file"
    else:
        package = Presentation()
        slide = package.slides.add_slide(package.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Uploaded office content"
        reader = PPTXReader(chunk=False)
        expected_name = "pptx_file"
    package.save(buffer)
    return buffer.getvalue(), reader, expected_name


@pytest.mark.parametrize("stream_kind", ["temporary", "spooled"])
@pytest.mark.parametrize("asynchronous", [False, True])
async def test_office_temporary_stream_names(office_package, stream_kind, asynchronous):
    payload, reader, fallback_name = office_package
    stream_factory = tempfile.TemporaryFile if stream_kind == "temporary" else tempfile.SpooledTemporaryFile
    with stream_factory(mode="w+b") as stream:
        stream.write(payload)
        stream.seek(0)
        documents = await reader.async_read(stream) if asynchronous else reader.read(stream)
        assert len(documents) == 1
        assert "Uploaded office content" in documents[0].content
        assert documents[0].name == (stream.name.split(".")[0] if isinstance(stream.name, str) else fallback_name)
        assert not stream.closed


@pytest.mark.parametrize("asynchronous", [False, True])
async def test_explicit_office_name_overrides_temporary_stream(office_package, asynchronous):
    payload, reader, _ = office_package
    with tempfile.TemporaryFile(mode="w+b") as stream:
        stream.write(payload)
        stream.seek(0)
        documents = (
            await reader.async_read(stream, name="Named upload")
            if asynchronous
            else reader.read(stream, name="Named upload")
        )
        assert len(documents) == 1
        assert documents[0].name == "Named upload"
        assert "Uploaded office content" in documents[0].content
